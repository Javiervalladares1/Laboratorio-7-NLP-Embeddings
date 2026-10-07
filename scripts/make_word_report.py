"""Informe sencillo en Word y figuras nuevas en blanco y negro.

Lee resultados existentes; no entrena ni vuelve a consultar el test.
Dependencias de documentos: python-docx, matplotlib, numpy.
"""
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/docx/Laboratorio_7_Informe.docx'
FIG = ROOT / 'figures/word'
MODELS = ['SGNS', 'gensim', 'GloVe']

def load(name):
    return json.loads((ROOT / 'results' / name).read_text())

def pct(x):
    return f'{100*x:.1f}' if x is not None else 'N/D'

def dec(x):
    return f'{x:.3f}' if x is not None else 'N/D'

def plots(runs, final):
    FIG.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({'font.size': 10, 'text.color': 'black',
        'axes.labelcolor': 'black', 'axes.edgecolor': 'black',
        'xtick.color': 'black', 'ytick.color': 'black',
        'axes.spines.top': False, 'axes.spines.right': False,
        'savefig.facecolor': 'white', 'figure.facecolor': 'white'})
    def save(name):
        plt.tight_layout(pad=.6)
        plt.savefig(FIG / name, dpi=190, bbox_inches='tight')
        plt.close()
    fig, axes = plt.subplots(1, 2, figsize=(9, 2.5))
    for name, style, marker in [('d100_n100', '-', 'o'), ('d100_w5', '--', 's'), ('d100_k5', ':', '^')]:
        run = next(r for r in runs if r['config']['name'] == name)
        for ax, key in zip(axes, ['loss', 'accuracy']):
            ys = [r['loss'] if key == 'loss' else 100*r['analogies']['accuracy']['total'] for r in run['records']]
            ax.plot([1, 2, 3], ys, color='black', linestyle=style, marker=marker, markersize=4, label=name)
        axes[0].legend(fontsize=8, frameon=False)
    axes[0].set(xlabel='Época', ylabel='Pérdida por par', title='Entrenamiento SGNS')
    axes[1].set(xlabel='Época', ylabel='Aciertos en analogías (%)', title='Evaluación por época')
    for ax in axes: ax.set_xticks([1, 2, 3])
    save('entrenamiento.png')
    size = sorted([r for r in runs if r['config']['dim']==100 and r['config']['window']==2 and r['config']['negative']==3], key=lambda r:r['corpus_tokens_known'])
    plt.figure(figsize=(6.4, 2.4))
    plt.plot([r['corpus_tokens_known']/1e6 for r in size], [100*r['records'][r['best_epoch']-1]['analogies']['accuracy']['total'] for r in size], 'o-', color='black', label='SGNS 100d, ventana 2, negativos 3')
    plt.axhline(100*final['GloVe']['analogies']['add']['accuracy']['total'], color='black', linestyle='--', label='GloVe 100d: 6,000 millones de tokens')
    plt.xlabel('Tokens WikiText conocidos (millones)'); plt.ylabel('Aciertos (%)')
    plt.legend(fontsize=8, frameon=False); save('corpus.png')
    clf = load('classification_final.json')
    plt.figure(figsize=(6.4, 2.8))
    for name, style, marker in [('random','-','o'), ('SGNS','--','s'), ('gensim',':','^'), ('GloVe','-.','D'), ('TFIDF',(0,(5,1,1,1)),'x')]:
        rows = sorted([r for r in clf if r['initialization']==name], key=lambda r:r['fraction'])
        plt.plot([100*r['fraction'] for r in rows], [r['test']['f1_macro'] for r in rows], color='black', linestyle=style, marker=marker, markersize=4, label=name)
    plt.xlabel('Datos de entrenamiento con etiqueta (%)'); plt.ylabel('F1 macro en test')
    plt.legend(ncol=3, fontsize=8, frameon=False); save('clasificacion.png')

def build():
    prep = load('preprocessing.json'); emb = load('embedding_final.json')
    sel = load('selection.json'); ref = load('gensim_training.json')
    clf = load('classification_final.json'); oov = load('classification_oov.json')
    full = {r['initialization']:r for r in clf if r['fraction']==1}
    small = {r['initialization']:r for r in clf if r['fraction']==.01}
    runs = [load(p.name) for p in sorted((ROOT/'results').glob('sgns_*.json'))]
    selected = next(r for r in runs if r['config']['name']==sel['best'])
    best = selected['records'][selected['best_epoch']-1]
    plots(runs, emb)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    sec.footer_distance = Inches(.3)
    for style in doc.styles:
        if style.type in (1,2):
            style.font.name = 'Arial'
            style.font.color.rgb = RGBColor(0,0,0)
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    norm = doc.styles['Normal']
    norm.font.size = Pt(10)
    norm.font.bold = False
    norm.paragraph_format.space_after = Pt(5)
    norm.paragraph_format.line_spacing = 1.05
    for name, size in [('Title',16), ('Heading 1',12), ('Heading 2',10.5)]:
        st = doc.styles[name]; st.font.size = Pt(size); st.font.bold = True
        st.paragraph_format.space_before = Pt(8 if name!='Title' else 0)
        st.paragraph_format.space_after = Pt(5)
    doc.styles['Caption'].font.size = Pt(8.5)
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].font.bold = False
    doc.core_properties.title = 'Laboratorio 7 NLP y embeddings'
    doc.core_properties.author = 'Javier Valladares 23045; Ian Cumes 23236'
    doc.core_properties.subject = 'Informe sencillo con resultados ejecutados'
    foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run('Laboratorio 7 · Valladares 23045 y Cumes 23236 · Página ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)

    def p(text, small=False):
        para = doc.add_paragraph(text)
        if small:
            for r in para.runs: r.font.size = Pt(9)
        return para
    def h(text): doc.add_heading(text, level=1)
    def table(rows, widths=None, size=9):
        t = doc.add_table(rows=0, cols=len(rows[0])); t.autofit=False
        if widths:
            for col, w in zip(t.columns,widths): col.width=Inches(w)
        props=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
        props.append(borders)
        for i,row in enumerate(rows):
            cells=t.add_row().cells
            if widths:
                for c,w in zip(cells,widths):c.width=Inches(w)
            trPr=t.rows[-1]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if i==0: trPr.append(OxmlElement('w:tblHeader'))
            for c,v in zip(cells,row):
                tcPr=c._tc.get_or_add_tcPr()
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'FFFFFF');tcPr.append(shd)
                c.text=str(v)
                par=c.paragraphs[0]
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2)
                par.paragraph_format.line_spacing=1
                for r in par.runs:
                    r.font.size=Pt(size);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return t
    def image(name,width,caption):
        par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(FIG/name),width=Inches(width))
        doc.add_paragraph(caption,style='Caption')
    def page():doc.add_page_break()

    # 1. Preparación y conceptos, sin portada independiente.
    doc.add_paragraph('Laboratorio 7 NLP y embeddings',style='Title')
    p('Javier Valladares 23045 e Ian Cumes 23236\nCC3092 · Deep Learning y Sistemas Inteligentes · Octubre de 2026')
    h('1. Objetivo')
    p('Entrenamos un modelo de palabras en PyTorch y lo comparamos con gensim y GloVe. Revisamos si los vectores representan relaciones entre palabras y si ayudan a clasificar noticias. Todos los resultados provienen de ejecuciones guardadas.')
    h('2. Preparación de los datos')
    train=prep['splits']['train']
    table([['Dato de WikiText 103','Cantidad'],['Artículos de entrenamiento',f"{train['articles']:,}"],['Líneas, incluidas las vacías',f"{train['lines']:,}"],['Tokens después de limpiar',f"{train['tokens']:,}"],['Tokens de validación / test',f"{prep['splits']['validation']['tokens']:,} / {prep['splits']['test']['tokens']:,}"],['Tokens usados / conocidos','22,000,000 / 21,336,925']], [4.7,2.4])
    p('Un token es una unidad del texto. Usamos sólo el inicio del conjunto de entrenamiento. Pasamos el texto a minúsculas, normalizamos Unicode, quitamos marcadores de WikiText y separamos puntuación, guiones y apóstrofos. Conservamos los números y las palabras comunes; no aplicamos stemming ni lematización.')
    p("Comparamos el tokenizador con NLTK en 12 oraciones. Por ejemplo, nuestra función separa can't en can y t; NLTK Treebank produce ca y n't. Estas diferencias afectan el vocabulario. La función conserva secuencias de letras a–z y números, por lo que algunas palabras con acentos se fragmentan.")
    table([['Frecuencia mínima','Palabras en vocabulario','Tokens fuera del vocabulario (%)']]+[[r['min_count'],f"{r['vocab_size']:,}",f"{r['oov_percent']:.2f}"] for r in prep['thresholds']], [1.6,2.5,3])
    cov=prep['coverage_full_train']
    p(f"Elegimos frecuencia mínima 20: quedan 38,569 palabras. La distribución se aproxima a Zipf: pocas palabras aparecen muchas veces. Las 10, 1,000 y 30,000 más frecuentes cubren {cov['10']:.2f}%, {cov['1000']:.2f}% y {cov['30000']:.2f}% del corpus completo. La pendiente medida es {prep['zipf_slope_rank_100_10000']:.2f}.")
    h('3. Conceptos utilizados')
    p('Embedding asigna un vector a cada palabra. from_pretrained carga vectores existentes y freeze decide si se actualizan. EmbeddingBag reúne los vectores de un texto por promedio, suma o máximo; offsets marca dónde empieza cada texto. padding_idx reserva el relleno y sparse permite gradientes dispersos.')
    p('CBOW predice una palabra a partir de su contexto. Skip gram predice el contexto a partir de una palabra. SGNS usa dos tablas: una para la palabra central y otra para el contexto; al terminar conservamos la primera. Word2Vec aprende a predecir contextos; GloVe aprende de conteos globales de coocurrencia.')
    page()

    # 2. Configuraciones, selección y dos gráficos clave.
    h('4. Entrenamiento del modelo propio')
    p('Implementamos SGNS en PyTorch con ejemplos negativos. Los negativos se muestrean según la frecuencia elevada a 0.75. Usamos 3 épocas, lotes de 1,024, SGD con tasa de aprendizaje de 0.025 a 0.001 y semilla 23045. Las ventanas no cruzan líneas. Equipo: Apple M1 Pro, 16 GB de RAM, CPU con 4 hilos; no se utilizó GPU.')
    rows=[['Configuración','Tokens M','Dim.','Ventana','Neg.','WS353','Aciertos %','Min.']]
    for r in runs:
        c=r['config'];e=r['records'][r['best_epoch']-1]
        rows.append([c['name'],f"{r['corpus_tokens_known']/1e6:.2f}",c['dim'],c['window'],c['negative'],dec(e['wordsim']['spearman']),pct(e['analogies']['accuracy']['total']),f"{r['training_seconds']/60:.1f}"])
    table(rows,[1.35,.75,.5,.7,.5,.7,.9,.7],8.5)
    p('M = millones de tokens conocidos; Dim. = tamaño del vector; Neg. = negativos por par; WS353 = correlación de Spearman en WordSim 353. Las métricas de la tabla corresponden a la mejor época, que fue la tercera en todas las configuraciones.',True)
    p(f"Seleccionamos {sel['best']}: 100 dimensiones, ventana 5 y 3 negativos. El criterio fue sumar WS353 y la proporción de aciertos en analogías, entre modelos con al menos 20 millones de tokens conocidos. SimLex se dejó para la evaluación final.")
    s=best['subsampling']
    p(f"Reducimos la aparición de palabras muy frecuentes con subsampling (umbral 0.00001). Se descartó {s['the']['discarded_percent']:.2f}% de the, {s['of']['discarded_percent']:.2f}% de of y {s['and']['discarded_percent']:.2f}% de and. En la última época, los pares bajaron de {best['pairs_before_subsampling']:,} a {best['pairs_after_subsampling']:,}.")
    image('entrenamiento.png',6.7,'La pérdida baja y los aciertos aumentan en estas tres configuraciones. Cambiar los negativos también cambia la escala de la pérdida.')
    image('corpus.png',5.6,'Con la misma configuración, un corpus mayor mejora las analogías. GloVe se entrenó con un corpus mucho mayor.')
    p(f"La referencia gensim usó el mismo corpus, vocabulario y parámetros equivalentes, y tardó {ref['training_seconds']/60:.1f} minutos frente a {selected['training_seconds']/60:.1f} del modelo propio. No esperamos vectores idénticos: gensim actualiza por pares con varios workers y PyTorch trabaja por lotes.")
    page()

    # 3. Comparación de analogías y explicación sencilla de la aritmética.
    h('5. Relaciones entre palabras')
    ev=emb['SGNS']['analogies']['add']
    p(f"Evaluamos {ev['evaluated']:,} de {ev['total_questions']:,} preguntas, con los mismos 30,000 candidatos para los tres modelos. La cobertura fue {100*ev['coverage']:.1f}%. 3CosAdd busca palabras cercanas al vector b − a + c; 3CosMul combina similitudes de forma multiplicativa. Ambos excluyen las palabras de la pregunta.")
    table([['Método y tipo de analogía','SGNS %','gensim %','GloVe %']]+[[method+' '+label,*[pct(emb[n]['analogies'][key]['accuracy'][kind]) for n in MODELS]] for method,key in [('3CosAdd','add'),('3CosMul','mul')] for label,kind in [('semántica','semantic'),('sintáctica','syntactic'),('total','total')]], [3.2,1.3,1.3,1.3])
    p('Resultados por categoría: cada celda muestra 3CosAdd / 3CosMul, en porcentaje. Se evaluaron las mismas preguntas en cada modelo.',True)
    labels=['Capitales comunes','Capitales del mundo','Monedas','Ciudades y estados','Familia','Adjetivo a adverbio','Opuestos','Comparativos','Superlativos','Participio presente','Gentilicios','Pasado','Plural','Plural de verbos']
    rows=[['Categoría','Preguntas','SGNS','gensim','GloVe']]
    for i,cat in enumerate(ev['categories']):
        rows.append([labels[i],cat['evaluated'],*[pct(emb[n]['analogies']['add']['categories'][i]['accuracy'])+' / '+pct(emb[n]['analogies']['mul']['categories'][i]['accuracy']) for n in MODELS]])
    table(rows,[2.3,.8,1.3,1.3,1.4],8.5)
    p('Además, probamos 21 analogías propias de 7 tipos: género, capital, gentilicio, comparativo, superlativo, pasado y plural. Nuestra implementación de 3CosAdd coincidió con gensim en todas las consultas cubiertas. El notebook incluye las cinco respuestas más cercanas y sus similitudes.')
    p('Ejemplo: man : woman :: king : queen. El resultado esperado no siempre queda en primer lugar.')
    rows=[['Modelo','Primera respuesta','Rango de queen','Coseno de queen']]
    for n in MODELS:
        q=emb[n]['individual']['queries'][0]
        rows.append([n,q['top5'][0][0],q['expected_rank'],dec(q['expected_cosine'])])
    table(rows,[1.2,2.3,1.8,1.8])
    p('Sin excluir las palabras de la pregunta, king fue la primera respuesta en los tres modelos. Por eso, esta operación representa una relación aproximada, no una igualdad exacta. El coseno medio entre diferencias de vectores fue 0.392 en SGNS, 0.381 en gensim y 0.619 en GloVe.')
    p('La visualización t-SNE del notebook contiene 500 palabras de países, números, animales, verbos y objetos. Países y números se agrupan con más claridad; los otros grupos se superponen. Una proyección a dos dimensiones no conserva todas las relaciones originales.',True)
    page()

    # 4. Clasificación: selección sólo en validación y curva de datos.
    h('6. Clasificación de noticias AG News')
    p('AG News tiene cuatro clases: mundo, deportes, negocios y ciencia/tecnología. Separamos 108,000 noticias para entrenamiento y 12,000 para validación; usamos las 7,600 noticias del test oficial sólo al final. La selección de época y de variante se hizo con F1 macro de validación.')
    p('El clasificador promedia los vectores con EmbeddingBag y usa una capa de 128 unidades, ReLU, dropout 0.2 y cuatro salidas. Entrenamos 10 épocas con lotes de 512 y tasa 0.003: Adam para el clasificador y SparseAdam para los embeddings. Comparamos vectores aleatorios, SGNS, gensim y GloVe, todos de 100 dimensiones. También usamos TF-IDF con regresión logística.')
    p('Congelado significa que los vectores no cambian. Ajustado significa que se actualizan con las noticias. F1 macro promedia el desempeño de las cuatro clases.')
    rows=[['Modelo','F1 val congelado','F1 val ajustado','Tiempo ajustado s']]
    for n in ['random','SGNS','gensim','GloVe']:
        ft=load('clf_'+n+'_finetune_p100.json')
        frozen=dec(load('clf_'+n+'_frozen_p100.json')['validation']['f1_macro']) if n!='random' else 'No aplica'
        rows.append([n,frozen,dec(ft['validation']['f1_macro']),f"{ft['training_seconds']:.1f}"])
    tf=load('clf_TFIDF_p100.json')
    rows.append(['TF-IDF','No aplica',dec(tf['validation']['f1_macro']),f"{tf['training_seconds']:.1f}"])
    table(rows,[1.4,1.9,1.9,1.9])
    p('Cada modelo con embeddings tiene 4.26 millones de parámetros: al congelar se entrenan sólo 13,444; al ajustar se entrenan todos. TF-IDF tiene 161,472 parámetros. Usa unigramas, frecuencia mínima 2, hasta 60,000 atributos y C = 4; el ajuste se hizo en seis bloques de diez pasos L-BFGS.',True)
    rows=[['Modelo elegido','Exactitud','Precisión macro','Recall macro','F1 macro']]
    for n in ['random','SGNS','gensim','GloVe','TFIDF']:
        t=full[n]['test'];rows.append([n+' ajustado' if n!='TFIDF' else 'TF-IDF',*[dec(t[k]) for k in ['accuracy','precision_macro','recall_macro','f1_macro']]])
    table(rows,[1.9,1.1,1.5,1.3,1.3])
    p('La tabla anterior usa todo el entrenamiento y muestra el test de las variantes elegidas en validación. Todas las fuentes de embeddings eligieron ajuste. GloVe obtuvo el mayor F1, aunque su diferencia frente a los vectores aleatorios fue pequeña.')
    image('clasificacion.png',5.5,'Cada punto corresponde a un modelo entrenado con 1%, 10%, 50% o 100% de las etiquetas. Se eligió la variante en validación y luego se evaluó una vez en test.')
    p(f"Con 1% de las etiquetas, F1 fue {dec(small['random']['test']['f1_macro'])} con vectores aleatorios, {dec(small['SGNS']['test']['f1_macro'])} con SGNS, {dec(small['GloVe']['test']['f1_macro'])} con GloVe y {dec(small['TFIDF']['test']['f1_macro'])} con TF-IDF. Los vectores preentrenados ayudan más cuando hay pocos ejemplos etiquetados.")
    p('Las curvas de pérdida y las matrices de confusión completas están en el notebook. Una confusión frecuente ocurre entre negocios y ciencia/tecnología, que comparten términos.',True)
    page()

    # 5. Tabla final y discusión con lenguaje directo.
    h('7. Comparación final')
    rows=[['Métrica','SGNS','gensim','GloVe'],['Tokens de entrenamiento','21.34 millones','21.34 millones','6,000 millones'],['Vocabulario / dimensiones','38,569 / 100','38,569 / 100','400,000 / 100'],['Tiempo / equipo',f"{selected['training_seconds']/60:.1f} min / M1 Pro CPU",f"{ref['training_seconds']/60:.1f} min / M1 Pro CPU",'No documentado para el archivo 100d']]
    rows.append(['WordSim 353 (Spearman)',*[dec(emb[n]['wordsim']['spearman']) for n in MODELS]])
    rows.append(['SimLex 999 (Spearman)',*[dec(emb[n]['simlex']['spearman']) for n in MODELS]])
    rows.append(['Pares evaluados WS / SimLex',*[f"{emb[n]['wordsim']['evaluated']} / {emb[n]['simlex']['evaluated']}" for n in MODELS]])
    rows.append(['Analogías 3CosAdd total (%)',*[pct(emb[n]['analogies']['add']['accuracy']['total']) for n in MODELS]])
    rows.append(['Analogías 3CosMul total (%)',*[pct(emb[n]['analogies']['mul']['accuracy']['total']) for n in MODELS]])
    rows.append(['Coseno de diferencias',*[dec(emb[n]['individual']['parallelism_mean']) for n in MODELS]])
    rows.append(['OOV AG News train / test (%)',*[f"{oov[n]['train']['source_oov_percent']:.2f} / {oov[n]['test']['source_oov_percent']:.2f}" for n in MODELS]])
    rows.append(['F1 test con 100% / 1% de etiquetas',*[dec(full[n]['test']['f1_macro'])+' / '+dec(small[n]['test']['f1_macro']) for n in MODELS]])
    table(rows,[2.65,1.45,1.45,1.55],8.5)
    p('OOV indica palabras que no existen en los vectores de origen. WordSim evalúa relaciones amplias entre palabras; SimLex se enfoca en similitud. Sus resultados no son equivalentes. Además, GloVe cubre más pares, por lo que estas correlaciones no usan exactamente los mismos ejemplos.',True)
    h('8. Discusión y conclusiones')
    p('El modelo propio aprendió relaciones útiles y obtuvo resultados cercanos a gensim. Aumentar la ventana de 2 a 5 elevó los aciertos de 11.8% a 23.4%; aumentar los negativos de 3 a 5, con ventana 2, los elevó a 15.3%. Aumentar la dimensión de 100 a 300 no mejoró los resultados en estas tres épocas.')
    p('SGNS resolvió mejor los gentilicios y peor las monedas. Sus aciertos fueron mayores en relaciones sintácticas que semánticas. Una pérdida menor no garantiza una buena analogía, porque el entrenamiento aprende contextos locales y la evaluación mide otra tarea.')
    p('GloVe fue mucho mejor en analogías, pero utilizó unas 281 veces más tokens. El experimento con prefijos de 25%, 50% y 100% muestra que más datos ayudan. No podemos atribuir toda la diferencia al tamaño: también cambian el corpus, el método y la optimización.')
    p('En clasificación, GloVe obtuvo F1 de 0.918 y los vectores aleatorios 0.917 con todas las etiquetas. Con pocas etiquetas, la ventaja de los vectores preentrenados fue mayor. Por eso, una mejora grande en analogías no implica una mejora igual de grande al clasificar noticias.')
    p('Estos resultados usan una sola semilla y tres épocas para embeddings, sin intervalos de confianza. Los prefijos pueden tener temas distintos. El vocabulario de AG News se obtuvo de todo train incluso en experimentos con pocas etiquetas. Gensim emitió avisos de Cython; la auditoría verificó que sus vectores eran finitos. Estos límites deben considerarse al interpretar las diferencias.',True)
    h('9. Repositorio y fuentes')
    p('Repositorio público: '+load('delivery.json')['repository'],True)
    p('El repositorio contiene el notebook ejecutado, código, resultados completos, gráficos y los vectores finales comprimidos. Las tablas de este informe se generan desde los archivos de resultados.',True)
    p('Fuentes: documentación de PyTorch (Embedding y EmbeddingBag); documentación de gensim (Word2Vec y KeyedVectors); Mikolov et al. (2013), arxiv.org/abs/1310.4546; Pennington et al. (2014), aclanthology.org/D14-1162; GloVe, nlp.stanford.edu/projects/glove; datasets Salesforce/wikitext y fancyzhx/ag_news en Hugging Face.',True)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__=='__main__':build()
